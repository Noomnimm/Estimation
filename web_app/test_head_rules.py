import unittest
from io import BytesIO
import pandas as pd
from openpyxl import load_workbook
from web_app.material_logic import (
    MaterialWorkbook, classify_wire_head, wire_head_multiplier, add_wire_materials,
    insulator_rate, SIZE_COL, HEAD_COL, MATERIAL_COL, CODE_COL, QTY_COL, TOTAL_COL,
)


class HeadRuleTests(unittest.TestCase):
    def equipment(self, head, wire1='185 SAC', wire2='185 SAC'):
        totals = {}
        add_wire_materials(totals, classify_wire_head(head), wire1, wire2, wire_head_multiplier(head))
        return {r[CODE_COL]: r[TOTAL_COL] for r in totals.values()}

    def test_double_heads(self):
        for head, preform in [('2BA.st 4.5 m', 6), ('2DE.st4.5', 6), ('2DDE.st 4.5m', 12)]:
            with self.subTest(head=head):
                values = self.equipment(head)
                self.assertEqual(values['1020260205'], preform)
                self.assertEqual(values['1030140011'], preform)
        values = self.equipment('2DDE.st 4.5m')
        for code in ['1020410027', '1020180001', '1020180008']:
            self.assertEqual(values[code], 6)
        self.assertEqual(self.equipment('2BA')['1020300103'], 12)
        small = self.equipment('2BA', wire2='50 SAC')
        self.assertEqual(small['1020260202'], 6)
        self.assertEqual(small['1020330104'], 6)
        self.assertEqual(small['1020330006'], 6)

    def test_single_phase(self):
        for head, expected in [('BA 1-P', 2), ('BA.AL 1P', 2), ('DE.CON 1-P', 2), ('DDE 1-P', 4), ('DDE.BL 1-P', 4)]:
            self.assertEqual(self.equipment(head)['1020260205'], expected)
        self.assertEqual(self.equipment('DDE 1-P')['1020410027'], 2)
        self.assertNotIn('1020410027', self.equipment('DDE.BL 1-P'))
        self.assertEqual(self.equipment('BA 1-P')['1020300103'], 4)

    def test_strain_and_pending(self):
        values = self.equipment('2DE', wire1='185 A')
        self.assertEqual(values['1030110004'], 6)
        self.assertNotIn('1030140011', values)
        for head in ['2BA.st 4.5m+DE.CON', '2DE.st4.5 + DE.CON']:
            self.assertIsNone(classify_wire_head(head))
        self.assertEqual(classify_wire_head('BA.SLK บน'), 'ba')  # Preserve existing wire rule.
        self.assertEqual(self.equipment('DDE.st 3m, LAT.SLK')['1020260205'], 6)

    def test_base_not_doubled(self):
        workbook = MaterialWorkbook()
        workbook.base_df = pd.DataFrame([{SIZE_COL: '99', HEAD_COL: '2DE.st4.5', MATERIAL_COL: 'Base', CODE_COL: 'SetExample', QTY_COL: 1}])
        row = {'size': '99', 'head': '2DE.st4.5', 'count': '1+1', 'wire1': '185 SAC'}
        result = workbook.calculate([[row]])
        values = {r[CODE_COL]: r[TOTAL_COL] for r in result['items']}
        self.assertEqual(values['SetExample'], 2)
        self.assertEqual(values['1020260205'], 12)

    def test_lat_as_de(self):
        for head in ['LAT.SLK บน', 'LAT.SLK ล่าง']:
            values = self.equipment(head, wire1='50 SAC')
            self.assertEqual(values['1020260202'], 3)
            self.assertEqual(values['1030140011'], 3)
            self.assertNotIn('1020410027', values)
            self.assertNotIn('1020180001', values)
            strain = self.equipment(head, wire1='185 A')
            self.assertEqual(strain['1030110004'], 3)
            self.assertNotIn('1030140011', strain)

    def test_mixed_lat(self):
        workbook = MaterialWorkbook()
        head = 'DDE.st 3m, LAT.SLK'
        workbook.base_df = pd.DataFrame([{SIZE_COL: '99', HEAD_COL: head, MATERIAL_COL: 'Base', CODE_COL: 'SetExample', QTY_COL: 1}])
        row = {'size': '99', 'head': head, 'count': '1+1', 'wire1': '185 SAC', 'wire2': '185 SAC'}
        with self.assertRaisesRegex(ValueError, 'LAT.SLK'):
            workbook.calculate([[row]])
        row['latWire'] = '50 SAC'
        result = workbook.calculate([[row]])
        values = {r[CODE_COL]: r[TOTAL_COL] for r in result['items']}
        self.assertEqual(values['SetExample'], 2)
        self.assertEqual(values['1020260205'], 12)
        self.assertEqual(values['1020260202'], 6)
        self.assertEqual(values['1030140011'], 18)
        self.assertEqual(values['1020410027'], 6)
        self.assertEqual(values['1020180001'], 6)

    def test_new_steel_combined_heads(self):
        cases = {
            'SP,DDE.BL st.4.5m': 'dde_bl',
            'DP,DDE.BL st.4.5m': 'dde_bl',
            'DP,DDE st.4.5m': 'dde',
            'DP,DE st.4.5m': 'de',
        }
        for head, expected_kind in cases.items():
            with self.subTest(head=head):
                self.assertEqual(classify_wire_head(head), expected_kind)
                values = self.equipment(head)
                expected_preform = 3 if expected_kind == 'de' else 6
                self.assertEqual(values['1020260205'], expected_preform)
                self.assertEqual(values['1030140011'], expected_preform)
                if expected_kind == 'dde':
                    self.assertEqual(values['1020410027'], 3)
                    self.assertEqual(values['1020180001'], 3)
                    self.assertEqual(values['1020180008'], 3)
                else:
                    self.assertNotIn('1020410027', values)

    def test_page_hardware_export(self):
        workbook = MaterialWorkbook()
        workbook.base_df = pd.DataFrame([{
            SIZE_COL: 'Dis & SF6', HEAD_COL: 'SF6', MATERIAL_COL: 'SF6 SET', CODE_COL: 'set24009', QTY_COL: 1,
        }])
        workbook.set_df = pd.DataFrame([
            {"Set": 'set24009', CODE_COL: '1020330006', "คำอธิบาย": 'HOTLINE BAIL-CLAMP,MAIN 70-185 SQ.MM.', "ติดตั้ง": 3},
            {"Set": 'set24009', CODE_COL: '1020330104', "คำอธิบาย": 'HOTLINE CLAMP,MAIN35-185,TAP50-185SQ.MM.', "ติดตั้ง": 3},
        ])
        pages = [[{
            'size': '14.3', 'head': 'DP,DDE st.4.5m', 'count': '1+1',
            'wire1': '185 SAC', 'wire2': '185 SAC',
        }, {
            'size': 'Dis & SF6', 'head': 'SF6', 'count': '1',
        }], [{
            'size': '12.2', 'head': 'DP,DE st.4.5m', 'count': '1',
            'wire1': '50 SAC',
        }]]
        exported = load_workbook(BytesIO(workbook.export_page_hardware(pages)), data_only=True)
        sheet = exported['ลูกถ้วยและอุปกรณ์']
        values = {(row[0], row[1], str(row[2] or '')): (row[3], row[4]) for row in sheet.iter_rows(min_row=4, values_only=True)}
        self.assertEqual(values[('ลูกถ้วย', 'ลูกถ้วยตั้ง', '')], (24, 6))
        self.assertEqual(values[('ลูกถ้วย', 'ลูกถ้วยนอน', '')], (48, 12))
        self.assertEqual(values[('อุปกรณ์ยึดสาย', 'PREFORMED D/E,SAC 22kV 185sq.mm. 29.78mm', '1020260205')], (12, None))
        self.assertEqual(values[('อุปกรณ์ยึดสาย', 'PREFORMED D/E,SAC 22kV 50sq.mm. 21.80mm', '1020260202')], (None, 3))
        self.assertEqual(values[('อุปกรณ์ยึดสาย', 'CLEVIS,THIMBLE,FOR PREFORMED DEAD-END', '1030140011')], (12, 3))
        self.assertEqual(values[('อุปกรณ์ยึดสาย', 'CONNECTOR,SPLICE,COMPRESSION TYPE,TENSIONLESS AL 185 SQ.MM.', '1020410027')], (6, None))
        self.assertEqual(values[('อุปกรณ์ยึดสาย', 'HOTLINE BAIL-CLAMP,MAIN 70-185 SQ.MM.', '1020330006')], (3, None))
        self.assertEqual(values[('อุปกรณ์ยึดสาย', 'HOTLINE CLAMP,MAIN35-185,TAP50-185SQ.MM.', '1020330104')], (3, None))
        self.assertEqual(sheet.freeze_panes, 'D4')
        detail_sheet = exported['ที่มารายการ']
        sf6_rows = [row for row in detail_sheet.iter_rows(min_row=2, values_only=True) if row[1] == 'SF6']
        self.assertEqual(len(sf6_rows), 2)
        self.assertTrue(all(row[2] == 'set24009' and row[5] == 3 for row in sf6_rows))

    def test_python_insulator_rules_match_confirmed_cases(self):
        expected = {
            'BA.st4.5m': (4, 12), '2BA st.4.5m': (12, 24),
            'SP,DDE.BL st.4.5m': (3, 24), 'DP,DDE.BL st.4.5m': (6, 24),
            'DP,DDE st.4.5m': (12, 24), 'DP,DE st.4.5m': (6, 12),
            'DDE.st 3m, LAT.SLK': (12, 36), 'DE.CON 1-P': (4, 8),
        }
        for head, rate in expected.items():
            self.assertEqual(insulator_rate(head), rate)

    def test_sf6_set_hardware_reconciles_with_wire_totals(self):
        workbook = MaterialWorkbook()
        workbook.base_df = pd.DataFrame([{
            SIZE_COL: 'Dis & SF6', HEAD_COL: 'SF6', MATERIAL_COL: 'SF6 SET', CODE_COL: 'set24009', QTY_COL: 1,
        }])
        workbook.set_df = pd.DataFrame([
            {"Set": 'set24009', CODE_COL: '1020330006', "คำอธิบาย": 'HOTLINE BAIL-CLAMP,MAIN 70-185 SQ.MM.', "ติดตั้ง": 3},
            {"Set": 'set24009', CODE_COL: '1020330104', "คำอธิบาย": 'HOTLINE CLAMP,MAIN35-185,TAP50-185SQ.MM.', "ติดตั้ง": 3},
        ])
        pages = [[
            {'size': '14.3', 'head': 'BA', 'count': 12, 'wire1': '185 SAC', 'wire2': '50 SAC'},
            {'size': 'Dis & SF6', 'head': 'SF6', 'count': 1},
        ]]
        exported = load_workbook(BytesIO(workbook.export_page_hardware(pages)), data_only=True)
        rows = list(exported['ลูกถ้วยและอุปกรณ์'].iter_rows(min_row=4, values_only=True))
        by_code = {str(row[2] or ''): row[3] for row in rows}
        self.assertEqual(by_code['1020330006'], 39)
        self.assertEqual(by_code['1020330104'], 39)

    def test_page_insulator_export_shows_totals_and_sources(self):
        workbook = MaterialWorkbook()
        pages = [[
            {'size': '14.3', 'head': 'BA', 'count': '1+1'},
            {'size': '14.3', 'head': 'DE', 'count': 1},
        ], [
            {'size': '12.2', 'head': 'SP 1-P', 'count': 3},
        ]]
        exported = load_workbook(BytesIO(workbook.export_page_insulators(pages)), data_only=True)
        summary = list(exported['สรุปลูกถ้วยแยกหน้า'].iter_rows(min_row=2, values_only=True))
        self.assertEqual(summary[0], (1, 8, 36, 44))
        self.assertEqual(summary[1], (2, 6, 0, 6))
        self.assertEqual(summary[2], ('รวมทุกหน้า', 14, 36, 50))
        details = list(exported['ที่มาลูกถ้วย'].iter_rows(min_row=2, values_only=True))
        self.assertEqual(details[0], (1, '14.3', 'BA', 2, 4, 12, 8, 24))
        self.assertEqual(details[1], (1, '14.3', 'DE', 1, 0, 12, 0, 12))
        self.assertEqual(details[2], (2, '12.2', 'SP 1-P', 3, 2, 0, 6, 0))

    def test_page_crossarm_export_expands_sets_and_applies_adjustments(self):
        workbook = MaterialWorkbook()
        workbook.base_df = pd.DataFrame([
            {SIZE_COL: '14.3', HEAD_COL: 'BA', MATERIAL_COL: 'BA SET', CODE_COL: 'set-ba', QTY_COL: 1},
            {SIZE_COL: '14.3', HEAD_COL: 'BA', MATERIAL_COL: 'คอน 2.5 เมตร', CODE_COL: '1000110004', QTY_COL: 2},
            {SIZE_COL: '14.3', HEAD_COL: 'BA', MATERIAL_COL: 'คอน 2 เมตร', CODE_COL: '1000110003', QTY_COL: -2},
        ])
        workbook.set_df = pd.DataFrame([
            {'Set': 'set-ba', CODE_COL: '1000110003', 'คำอธิบาย': 'คอนคอนกรีต 2 เมตร', 'ติดตั้ง': 2},
            {'Set': 'set-ba', CODE_COL: '1010200001', 'คำอธิบาย': 'เหล็กค้ำคอน', 'ติดตั้ง': 4},
        ])
        pages = [[{'size': '14.3', 'head': 'BA', 'count': 3}]]
        exported = load_workbook(BytesIO(workbook.export_page_crossarms(pages)), data_only=True)
        summary = list(exported['สรุปคอนแยกหน้า'].iter_rows(min_row=2, values_only=True))
        self.assertEqual(summary, [('คอน 2.5 เมตร', '1000110004', 6, 6)])
        details = list(exported['ที่มาคอน'].iter_rows(min_row=2, values_only=True))
        self.assertEqual(details, [(1, '14.3', 'BA', 3, 'คอน 2.5 เมตร', '1000110004', 6)])


if __name__ == '__main__':
    unittest.main()
