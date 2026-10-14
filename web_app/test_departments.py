import unittest
from io import BytesIO
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

from web_app.material_logic import (
    CODE_COL,
    DEPARTMENT_COL,
    HEAD_COL,
    MATERIAL_COL,
    QTY_COL,
    SET_COL,
    SET_DESC_COL,
    SET_INSTALL_COL,
    SET_REUSE_COL,
    SIZE_COL,
    THREE_PHASE_SCAFFOLD_HEAD,
    TOTAL_COL,
    INSULATOR_UPRIGHT_COL,
    INSULATOR_HORIZONTAL_COL,
    MaterialWorkbook,
    three_phase_config_head,
    three_phase_scaffold_config_head,
)


class DepartmentTests(unittest.TestCase):
    def setUp(self):
        self.workbook = MaterialWorkbook()
        self.workbook.base_df = pd.DataFrame([
            {SIZE_COL: "1", HEAD_COL: "HEAD", MATERIAL_COL: "HV", CODE_COL: "100", QTY_COL: 1, DEPARTMENT_COL: "แผนกแรงสูง"},
            {SIZE_COL: "1", HEAD_COL: "HEAD", MATERIAL_COL: "TX", CODE_COL: "200", QTY_COL: 1, DEPARTMENT_COL: "แผนกหม้อแปลง", INSULATOR_UPRIGHT_COL: 7, INSULATOR_HORIZONTAL_COL: 11},
        ])

    def test_department_filters_selectors_and_calculation(self):
        self.assertIn("แผนกแรงต่ำ", self.workbook.get_departments())
        self.assertEqual(self.workbook.get_sizes("แผนกแรงสูง TAC"), [])
        self.assertEqual(self.workbook.get_heads("1", "แผนกหม้อแปลง"), ["HEAD"])
        result = self.workbook.calculate([[
            {"department": "แผนกหม้อแปลง", "size": "1", "head": "HEAD", "count": "3"},
        ]])
        self.assertEqual(result["items"], [{MATERIAL_COL: "TX", CODE_COL: "200", TOTAL_COL: 3.0}])
        self.assertEqual(self.workbook.get_insulator_rate("1", "HEAD", "แผนกหม้อแปลง"), (7.0, 11.0))

    def test_transmission_structure_selector_groups_ss_and_sd_by_catalog_code(self):
        self.workbook.base_df = pd.DataFrame([
            {SIZE_COL: "101", HEAD_COL: "SS-TG-2 ASSEMBLY NO. 5212A", MATERIAL_COL: "SS tangent", CODE_COL: "Set50202", QTY_COL: 1, DEPARTMENT_COL: "แผนกสายส่ง"},
            {SIZE_COL: "205", HEAD_COL: "SD-TG-3 ASSEMBLY NO. 5264A", MATERIAL_COL: "SD tangent", CODE_COL: "Set50303", QTY_COL: 1, DEPARTMENT_COL: "แผนกสายส่ง"},
            {SIZE_COL: "301", HEAD_COL: "DS-TG-1 ASSEMBLY NO. 5301", MATERIAL_COL: "DS tangent", CODE_COL: "Set53000", QTY_COL: 1, DEPARTMENT_COL: "แผนกสายส่ง"},
            {SIZE_COL: "1", HEAD_COL: "OTHER STRUCTURE", MATERIAL_COL: "Other", CODE_COL: "Set10001", QTY_COL: 1, DEPARTMENT_COL: "แผนกสายส่ง"},
            {SIZE_COL: "102", HEAD_COL: "SS-TG-3", MATERIAL_COL: "HV", CODE_COL: "Set50015", QTY_COL: 1, DEPARTMENT_COL: "แผนกแรงสูง"},
        ])

        self.assertEqual(self.workbook.get_transmission_structures(), [
            {"keycode": "101", "circuit": "SS", "type": "TG", "head": "SS-TG-2 ASSEMBLY NO. 5212A", "code": "Set50202"},
            {"keycode": "205", "circuit": "SD", "type": "TG", "head": "SD-TG-3 ASSEMBLY NO. 5264A", "code": "Set50303"},
            {"keycode": "301", "circuit": "DS", "type": "TG", "head": "DS-TG-1 ASSEMBLY NO. 5301", "code": "Set53000"},
        ])

    def test_combined_hardware_export_contains_insulator_sheets(self):
        data = self.workbook.export_page_hardware_combined([[
            {"department": "แผนกหม้อแปลง", "size": "1", "head": "HEAD", "count": "2"},
        ]])
        workbook = load_workbook(BytesIO(data), read_only=True)
        self.assertEqual(workbook.sheetnames, [
            "ลูกถ้วยและอุปกรณ์", "ที่มารายการ", "สรุปลูกถ้วยแยกหน้า", "ที่มาลูกถ้วย",
        ])

    def test_catalog_code_override_is_stored_as_ten_digit_text(self):
        raw = pd.DataFrame([[None] * 9 for _ in range(8)])
        raw.iloc[7, 1] = "15"
        raw.iloc[7, 2] = "หม้อแปลงขนาด 30 เควีเอ"
        raw.iloc[7, 8] = "1050010004"
        workbook = MaterialWorkbook()
        workbook.base_df = pd.DataFrame(columns=[SIZE_COL, HEAD_COL, MATERIAL_COL, CODE_COL, QTY_COL, DEPARTMENT_COL])
        with patch("web_app.material_logic.pd.read_excel", return_value=raw):
            workbook.load_keycode_catalog("unused.xlsx", "แผนกหม้อแปลง", {"15": "1050000011"})
        self.assertEqual(workbook.base_df.iloc[0][CODE_COL], "1050000011")

    def test_department_base_preserves_negative_quantities_and_normalizes_codes(self):
        source = pd.DataFrame([{
            "รายการ": "30kVA 1P", MATERIAL_COL: "อุปกรณ์ทดสอบ",
            CODE_COL: "1-05-000-0011", QTY_COL: -2,
        }])
        workbook = MaterialWorkbook()
        workbook.base_df = pd.DataFrame(columns=[SIZE_COL, HEAD_COL, MATERIAL_COL, CODE_COL, QTY_COL, DEPARTMENT_COL])
        with patch("web_app.material_logic.pd.read_excel", return_value=source):
            workbook.load_department_base("unused.xlsx", "แผนกหม้อแปลง", "หม้อแปลง")
        row = workbook.base_df.iloc[0]
        self.assertEqual(row[CODE_COL], "1050000011")
        self.assertEqual(row[QTY_COL], -2)

    def test_high_voltage_auxiliary_catalog_copies_tac_and_adds_ground_sets_idempotently(self):
        source_rows = [
            {SIZE_COL: "Surge & Fuse", HEAD_COL: "Drop-out Fuse", MATERIAL_COL: "CUT-OUT", CODE_COL: "1040010002", QTY_COL: 3, DEPARTMENT_COL: "แผนกแรงสูง"},
            {SIZE_COL: "Surge & Fuse", HEAD_COL: "LA 20-21kV 5kA", MATERIAL_COL: "LIGHTNING ARRESTER", CODE_COL: "1040000000", QTY_COL: 3, DEPARTMENT_COL: "แผนกแรงสูง"},
            {SIZE_COL: "Ground", HEAD_COL: "HV.GROUNDING GR-1", MATERIAL_COL: "HV.GROUNDING GR-1", CODE_COL: "Set20515", QTY_COL: 1, DEPARTMENT_COL: "แผนกแรงสูง"},
        ]
        workbook = MaterialWorkbook()
        workbook.base_df = pd.DataFrame(source_rows)

        workbook.ensure_high_voltage_auxiliary_catalog()
        workbook.ensure_high_voltage_auxiliary_catalog()

        for department in ("แผนกแรงสูง", "แผนกแรงสูง TAC"):
            ground = workbook.base_df[
                (workbook.base_df[SIZE_COL] == "Ground")
                & (workbook.base_df[DEPARTMENT_COL] == department)
            ]
            self.assertEqual(len(ground), 5)
            self.assertEqual(
                ground[CODE_COL].tolist(),
                ["Set20515", "Set20516", "Set20517", "Set20518", "Set20519"],
            )
            self.assertEqual(
                ground[HEAD_COL].tolist(),
                [
                    f"HV.GROUNDING GR-{number} (ASSEMBLY NO.9706, 9701B)"
                    for number in range(1, 6)
                ],
            )

        tac_surge = workbook.base_df[
            (workbook.base_df[SIZE_COL] == "Surge & Fuse")
            & (workbook.base_df[DEPARTMENT_COL] == "แผนกแรงสูง TAC")
        ]
        self.assertEqual(len(tac_surge), 2)
        self.assertEqual(
            set(tac_surge[CODE_COL]),
            {"1040010002", "1040000000"},
        )

    def test_three_phase_variants_keep_manual_surge_and_material_adjustment_rows(self):
        source = pd.DataFrame([
            {"รายการ": "30kVA 1P", MATERIAL_COL: "TR 1 เฟส", CODE_COL: "1050000002", QTY_COL: 1},
            {"รายการ": "30kVA 1P", MATERIAL_COL: "ชุดแรงสูง", CODE_COL: "Set40201", QTY_COL: 1},
            {"รายการ": "30kVA 1P", MATERIAL_COL: "สวิตช์", CODE_COL: "Set14141", QTY_COL: 1},
            {"รายการ": "30kVA 1P", MATERIAL_COL: "SA เก่า", CODE_COL: "1040000007", QTY_COL: -2},
            {"รายการ": "30kVA 1P", MATERIAL_COL: "SA ใหม่", CODE_COL: "1040000000", QTY_COL: 2},
            {"รายการ": "30kVA 1P", MATERIAL_COL: "ฟิวส์ LT", CODE_COL: "Set14013", QTY_COL: 1},
            {"รายการ": "30kVA 1P", MATERIAL_COL: "สาย LT 1 เฟส 2 สาย", CODE_COL: "Set40104", QTY_COL: 1},
            {"รายการ": "30kVA 1P", MATERIAL_COL: "BAIL 35-70", CODE_COL: "1020330005", QTY_COL: -2},
            {"รายการ": "30kVA 1P", MATERIAL_COL: "BAIL 70-185", CODE_COL: "1020330006", QTY_COL: 2},
            {"รายการ": "30kVA 1P", MATERIAL_COL: "BUSHING COVERS", CODE_COL: "1090250038", QTY_COL: 2},
            {"รายการ": "30kVA 1P", MATERIAL_COL: "DROPOUT COVERS", CODE_COL: "1090250040", QTY_COL: 2},
            {"รายการ": "30kVA 1P", MATERIAL_COL: "FUSE LINK", CODE_COL: "1040030000", QTY_COL: 2},
        ])
        workbook = MaterialWorkbook()
        workbook.base_df = pd.DataFrame(columns=[SIZE_COL, HEAD_COL, MATERIAL_COL, CODE_COL, QTY_COL, DEPARTMENT_COL])
        with patch("web_app.material_logic.pd.read_excel", return_value=source):
            workbook.load_department_base("unused.xlsx", "แผนกหม้อแปลง", "หม้อแปลง")

        head = three_phase_config_head(100, 1, 0)
        rows = workbook.base_df[(workbook.base_df[HEAD_COL] == head) & (workbook.base_df[DEPARTMENT_COL] == "แผนกหม้อแปลง")]
        rows_by_code = {row[CODE_COL]: row[QTY_COL] for _, row in rows.iterrows()}
        self.assertEqual(rows_by_code["1040000007"], -3)
        self.assertEqual(rows_by_code["1040000000"], 3)
        self.assertEqual(rows_by_code["1020330005"], -2)
        self.assertEqual(rows_by_code["1090250040"], 2)
        self.assertEqual(rows_by_code["Set40114"], 1)
        self.assertNotIn("Set40104", set(rows[CODE_COL]))

        two_lt_head = three_phase_config_head(100, 2, 0)
        two_lt_rows = workbook.base_df[(workbook.base_df[HEAD_COL] == two_lt_head) & (workbook.base_df[DEPARTMENT_COL] == "แผนกหม้อแปลง")]
        self.assertEqual(two_lt_rows.loc[two_lt_rows[CODE_COL] == "Set40114", QTY_COL].iloc[0], 2)

        fifty_kva_head = three_phase_config_head(50, 1, 0)
        fifty_kva_rows = workbook.base_df[(workbook.base_df[HEAD_COL] == fifty_kva_head) & (workbook.base_df[DEPARTMENT_COL] == "แผนกหม้อแปลง")]
        self.assertEqual(fifty_kva_rows.loc[fifty_kva_rows[CODE_COL] == "Set40112", QTY_COL].iloc[0], 1)

    def test_three_phase_scaffold_head_supports_one_or_two_lt_sides_and_optional_stub(self):
        source = pd.DataFrame([{
            "รายการ": "30kVA 1P", MATERIAL_COL: "TR 1 เฟส", CODE_COL: "1050000002", QTY_COL: 1,
        }])
        workbook = MaterialWorkbook()
        workbook.base_df = pd.DataFrame(columns=[SIZE_COL, HEAD_COL, MATERIAL_COL, CODE_COL, QTY_COL, DEPARTMENT_COL])
        with patch("web_app.material_logic.pd.read_excel", return_value=source):
            workbook.load_department_base("unused.xlsx", "แผนกหม้อแปลง", "หม้อแปลง")

        transformer_heads = workbook.get_heads("หม้อแปลง", "แผนกหม้อแปลง")
        self.assertIn("หม้อแปลง 3 เฟส 22 kV (Sealed)", transformer_heads)
        self.assertNotIn(THREE_PHASE_SCAFFOLD_HEAD, transformer_heads)
        integrated_scaffold_items = workbook.calculate([[
            {"department": "แผนกหม้อแปลง", "size": "หม้อแปลง", "head": "หม้อแปลง 3 เฟส 22 kV (Sealed)",
             "count": 1, "transformerKva": 100, "transformerType": "short-circuit",
             "transformerMounting": "scaffold", "transformerLtCircuits": 2,
             "transformerPlatformSides": 2, "transformerHasConcreteStub": True},
        ]])["items"]
        integrated_scaffold_codes = {row[CODE_COL]: row[TOTAL_COL] for row in integrated_scaffold_items}
        self.assertEqual(integrated_scaffold_codes["1050010052"], 1)
        self.assertEqual(integrated_scaffold_codes["Set14143"], 2)
        self.assertEqual(integrated_scaffold_codes["Set14021"], 2)
        self.assertEqual(integrated_scaffold_codes["Set14304"], 1)
        self.assertIn("set40211", workbook.transformer_scaffold_animal_barrier_sets)

        # Scaffold Surge always uses the crossarm position; only NGR and station
        # distance select the arrester variant. Keep the fixed default-tank
        # offset intact while replacing the three crossarm arresters.
        for ngr, within_3km, expected_code in [
            (False, False, "1040000000"),
            (False, True, "1040000001"),
            (True, False, "1040000002"),
            (True, True, "1040000003"),
        ]:
            surge_items = workbook.calculate([[
                {"department": "แผนกหม้อแปลง", "size": "หม้อแปลง", "head": "หม้อแปลง 3 เฟส 22 kV (Sealed)",
                 "count": 1, "transformerMounting": "scaffold", "transformerKva": 100,
                 "transformerLtCircuits": 1, "surgeNgr": ngr, "surgeWithin3km": within_3km,
                 "surgeMounting": "tank"},
            ]])["items"]
            surge_codes = {row[CODE_COL]: row[TOTAL_COL] for row in surge_items}
            self.assertIn(expected_code, surge_codes, surge_codes)
            self.assertEqual(surge_codes[expected_code], 3)
            self.assertEqual(surge_codes["1040000007"], -3)

        selectable_fuse_items = workbook.calculate([[
            {"department": "แผนกหม้อแปลง", "size": "หม้อแปลง", "head": "หม้อแปลง 3 เฟส 22 kV (Sealed)",
             "count": 1, "transformerKva": 100, "transformerMounting": "scaffold",
             "transformerLtCircuits": 2, "transformerScaffoldFuseChoice": 0},
        ]])["items"]
        selectable_fuse_codes = {row[CODE_COL]: row[TOTAL_COL] for row in selectable_fuse_items}
        self.assertEqual(selectable_fuse_codes["Set14020"], 1)
        self.assertEqual(selectable_fuse_codes["Set14022"], 1)
        self.assertNotIn("Set14021", selectable_fuse_codes)

        one_side_fuse_items = workbook.calculate([[
            {"department": "แผนกหม้อแปลง", "size": "หม้อแปลง", "head": "หม้อแปลง 3 เฟส 22 kV (Sealed)",
             "count": 1, "transformerKva": 100, "transformerMounting": "scaffold",
             "transformerLtCircuits": 1, "transformerScaffoldFuseChoice": 0},
        ]])["items"]
        one_side_fuse_codes = {row[CODE_COL]: row[TOTAL_COL] for row in one_side_fuse_items}
        self.assertEqual(one_side_fuse_codes["Set14023"], 1)

        selected_pole_items = workbook.calculate([[
            {"department": "แผนกหม้อแปลง", "size": "หม้อแปลง", "head": "หม้อแปลง 3 เฟส 22 kV (Sealed)",
             "count": 1, "transformerKva": 100, "transformerType": "short-circuit",
             "transformerMounting": "scaffold", "transformerLtCircuits": 1,
             "transformerScaffoldPoleCode": "1000010012", "transformerScaffoldStubCode": "Set14317"},
        ]])["items"]
        selected_pole_codes = {row[CODE_COL]: row[TOTAL_COL] for row in selected_pole_items}
        self.assertEqual(selected_pole_codes["1000010012"], -1)
        self.assertNotIn("1000010004", selected_pole_codes)
        self.assertEqual(selected_pole_codes["Set14317"], 1)
        self.assertNotIn("Set14304", selected_pole_codes)
        with self.assertRaisesRegex(ValueError, "ตอม่อที่เลือกไม่ตรงกับขนาดเสา"):
            workbook.calculate([[
                {"department": "แผนกหม้อแปลง", "size": "หม้อแปลง", "head": "หม้อแปลง 3 เฟส 22 kV (Sealed)",
                 "count": 1, "transformerKva": 100, "transformerMounting": "scaffold",
                 "transformerScaffoldPoleCode": "1000010012", "transformerScaffoldStubCode": "Set14304"},
            ]])

        one_side = three_phase_scaffold_config_head(1, False)
        one_side_items = workbook.calculate([[
            {"department": "แผนกหม้อแปลง", "size": "หม้อแปลง", "head": THREE_PHASE_SCAFFOLD_HEAD,
             "count": 1, "transformerPlatformSides": 1, "transformerHasConcreteStub": False},
        ]])["items"]
        one_side_codes = {row[CODE_COL]: row[TOTAL_COL] for row in one_side_items}
        self.assertEqual(one_side_codes["1050010052"], 1)
        self.assertEqual(one_side_codes["Set14143"], 1)
        self.assertEqual(one_side_codes["Set14023"], 1)
        self.assertNotIn("Set14021", one_side_codes)
        self.assertNotIn("Set14304", one_side_codes)
        self.assertEqual(one_side_codes["1010110402"], 12)
        self.assertEqual(one_side_codes["1010110401"], -12)
        for animal_barrier_code in ("1030010200", "1040010015", "1040010016", "1010100002", "1010230000"):
            self.assertNotIn(animal_barrier_code, one_side_codes)

        two_side = three_phase_scaffold_config_head(2, True)
        two_side_rows = workbook.base_df[
            (workbook.base_df[HEAD_COL] == two_side)
            & (workbook.base_df[DEPARTMENT_COL] == "แผนกหม้อแปลง")
        ]
        two_side_codes = two_side_rows.groupby(CODE_COL)[QTY_COL].sum().to_dict()
        self.assertEqual(two_side_codes["Set14143"], 2)
        self.assertEqual(two_side_codes["Set14021"], 2)
        self.assertEqual(two_side_codes["Set40114"], 2)
        self.assertNotIn("Set14023", two_side_codes)
        self.assertEqual(two_side_codes["Set14304"], 1)
        for animal_barrier_code in ("1030010200", "1040010015", "1040010016", "1010100002", "1010230000"):
            self.assertNotIn(animal_barrier_code, two_side_codes)

        demolition_items = workbook.calculate([
            [{"department": "แผนกหม้อแปลง", "size": "หม้อแปลง", "head": THREE_PHASE_SCAFFOLD_HEAD,
              "count": 1, "workType": "install", "transformerPlatformSides": 2,
              "transformerHasConcreteStub": True}],
            [{"department": "แผนกหม้อแปลง", "size": "หม้อแปลง", "head": THREE_PHASE_SCAFFOLD_HEAD,
              "count": 1, "workType": "demolition", "transformerPlatformSides": 2,
              "transformerHasConcreteStub": True}],
        ])["items"]
        demolition_codes = {row[CODE_COL]: row["จำนวนรื้อถอน"] for row in demolition_items}
        self.assertEqual(demolition_codes["Set14143"], 2)
        self.assertEqual(demolition_codes["Set14021"], 0)
        self.assertEqual(demolition_codes["1010110402"], 0)
        for animal_barrier_code in ("1030010200", "1040010015", "1040010016", "1010100002", "1010230000"):
            self.assertNotIn(animal_barrier_code, demolition_codes)

        workbook.set_df = pd.DataFrame([{
            SET_COL: "Set14304", CODE_COL: "1000020102", SET_DESC_COL: "STUB,PRESTRESSED CONCRETE",
            SET_INSTALL_COL: 1, SET_REUSE_COL: 1,
        }])
        workbook.calculate([
            [{"department": "แผนกหม้อแปลง", "size": "หม้อแปลง", "head": THREE_PHASE_SCAFFOLD_HEAD,
              "count": 1, "workType": "install", "transformerPlatformSides": 1,
              "transformerHasConcreteStub": True}],
            [{"department": "แผนกหม้อแปลง", "size": "หม้อแปลง", "head": THREE_PHASE_SCAFFOLD_HEAD,
              "count": 1, "workType": "demolition", "transformerPlatformSides": 1,
              "transformerHasConcreteStub": True}],
        ])
        self.assertEqual(workbook.get_set_components("Set14304", 1, "install")["items"][0]["quantity"], 1)
        self.assertEqual(workbook.get_set_components("Set14304", 1, "demolition")["items"], [])
        workbook.set_df = pd.DataFrame([{
            SET_COL: "Set14317", CODE_COL: "1000020401", SET_DESC_COL: "STUB,PRESTRESSED CONCRETE,0.36X0.36M.,5.50M.LONG",
            SET_INSTALL_COL: 1, SET_REUSE_COL: 1,
        }])
        workbook.calculate([[
            {"department": "แผนกหม้อแปลง", "size": "หม้อแปลง", "head": "หม้อแปลง 3 เฟส 22 kV (Sealed)",
             "count": 1, "workType": "demolition", "transformerKva": 100,
             "transformerMounting": "scaffold", "transformerScaffoldPoleCode": "1000010012",
             "transformerScaffoldStubCode": "Set14317"},
        ]])
        self.assertEqual(workbook.get_set_components("Set14317", 1, "demolition")["items"][0]["quantity"], 1)

    def test_scaffold_set40211_filters_animal_barrier_only_for_scaffold(self):
        barrier_codes = ("1030010200", "1040010015", "1040010016", "1010100002", "1010230000")
        workbook = MaterialWorkbook()
        workbook.set_df = pd.DataFrame([
            {SET_COL: "Set40211", CODE_COL: code, SET_DESC_COL: f"Barrier item {code}", SET_INSTALL_COL: 1, SET_REUSE_COL: 0}
            for code in barrier_codes
        ] + [
            {SET_COL: "Set40211", CODE_COL: "1000010004", SET_DESC_COL: "POLE", SET_INSTALL_COL: 1, SET_REUSE_COL: 1},
        ])
        workbook.transformer_scaffold_animal_barrier_sets = {"set40211"}
        scaffold_components = workbook.get_set_components("Set40211")["items"]
        self.assertEqual([row["code"] for row in scaffold_components], ["1000010004"])

        workbook.transformer_scaffold_animal_barrier_sets = set()
        ordinary_components = workbook.get_set_components("Set40211")["items"]
        self.assertEqual({row["code"] for row in ordinary_components}, {*barrier_codes, "1000010004"})

    def test_transformer_surge_arrester_follows_ngr_and_distance(self):
        workbook = MaterialWorkbook()
        workbook.base_df = pd.DataFrame([
            {SIZE_COL: "หม้อแปลง", HEAD_COL: "30kVA 1P", MATERIAL_COL: "OLD SA", CODE_COL: "1040000007", QTY_COL: -2, DEPARTMENT_COL: "แผนกหม้อแปลง"},
            {SIZE_COL: "หม้อแปลง", HEAD_COL: "30kVA 1P", MATERIAL_COL: "NEW LA", CODE_COL: "1040000000", QTY_COL: 2, DEPARTMENT_COL: "แผนกหม้อแปลง"},
        ])
        cases = [
            (False, False, "1040000007", "1040000000"),
            (False, True, "1040000007", "1040000001"),
            (True, False, "1040000008", "1040000002"),
            (True, True, "1040000008", "1040000003"),
        ]
        for ngr, within, removed_code, installed_code in cases:
            result = workbook.calculate([[
                {"department": "แผนกหม้อแปลง", "size": "หม้อแปลง", "head": "30kVA 1P", "count": 1,
                 "surgeNgr": ngr, "surgeWithin3km": within},
            ]])
            values = {row[CODE_COL]: row[TOTAL_COL] for row in result["items"]}
            self.assertEqual(values, {removed_code: -2.0, installed_code: 2.0})
        for ngr, tank_code in [(False, "1040000007"), (True, "1040000008")]:
            result = workbook.calculate([[
                {"department": "แผนกหม้อแปลง", "size": "หม้อแปลง", "head": "30kVA 1P", "count": 1,
                 "surgeNgr": ngr, "surgeWithin3km": False, "surgeMounting": "tank"},
            ]])
            self.assertEqual({row[CODE_COL]: row[TOTAL_COL] for row in result["items"]}, {tank_code: 2.0})

    def test_editable_three_phase_variant_keeps_selected_surge_and_set_adjustment_balanced(self):
        head = three_phase_config_head(100, 1, 0)
        workbook = MaterialWorkbook()
        workbook.base_df = pd.DataFrame([
            {SIZE_COL: "หม้อแปลง", HEAD_COL: head, MATERIAL_COL: "ชุดหม้อแปลง", CODE_COL: "Set40203", QTY_COL: 1, DEPARTMENT_COL: "แผนกหม้อแปลง"},
            {SIZE_COL: "หม้อแปลง", HEAD_COL: head, MATERIAL_COL: "S.A. default remove", CODE_COL: "1040000007", QTY_COL: -3, DEPARTMENT_COL: "แผนกหม้อแปลง"},
            {SIZE_COL: "หม้อแปลง", HEAD_COL: head, MATERIAL_COL: "S.A. selected", CODE_COL: "1040000000", QTY_COL: 3, DEPARTMENT_COL: "แผนกหม้อแปลง"},
        ])
        workbook.set_df = pd.DataFrame([
            {SET_COL: "Set40203", CODE_COL: "1040000007", SET_DESC_COL: "S.A. default", SET_INSTALL_COL: 3},
            {SET_COL: "Set40203", CODE_COL: "1050010201", SET_DESC_COL: "Transformer", SET_INSTALL_COL: 1},
        ])

        workbook.calculate([[
            {"department": "แผนกหม้อแปลง", "size": "หม้อแปลง", "head": head, "count": 1,
             "surgeNgr": False, "surgeWithin3km": False, "surgeMounting": "crossarm"},
        ]])
        result = workbook.expand_set()
        values = {row[CODE_COL]: row[TOTAL_COL] for row in result["items"]}

        self.assertEqual(values["1040000000"], 3)
        self.assertEqual(values["1040000007"], 0)

    def test_export_department_summaries_creates_one_sheet_per_department(self):
        data = self.workbook.export_department_summaries([
            {"name": "แผนกแรงสูง", "results": [{MATERIAL_COL: "BOLT", CODE_COL: "101", TOTAL_COL: 2}]},
            {"name": "แผนกหม้อแปลง", "results": [{MATERIAL_COL: "TR", CODE_COL: "202", TOTAL_COL: 1}]},
            {"name": "แผนกว่าง", "results": []},
        ])

        exported = load_workbook(BytesIO(data), data_only=True)
        self.assertEqual(exported.sheetnames, ["แผนกแรงสูง", "แผนกหม้อแปลง", "แผนกว่าง"])
        self.assertEqual(exported["แผนกแรงสูง"]["A2"].value, "BOLT")
        self.assertEqual(exported["แผนกหม้อแปลง"]["B2"].value, "202")
        self.assertEqual(exported["แผนกว่าง"]["A2"].value, "ยังไม่มีผลคำนวณในแผนกนี้")

    def test_page_exports_prefix_each_department_sheet(self):
        sample = BytesIO()
        pd.DataFrame([{"รายการ": "ข้อมูลทดสอบ"}]).to_excel(sample, index=False, sheet_name="รายงาน")
        sample_bytes = sample.getvalue()
        departments = [
            {"name": "แผนกแรงสูง", "pages": [[]]},
            {"name": "แผนกหม้อแปลง", "pages": [[]]},
        ]

        for export_name, method_name in (
            ("export_department_page_summary", "export_page_summary"),
            ("export_department_page_hardware", "export_page_hardware_combined"),
            ("export_department_page_crossarms", "export_page_crossarms"),
        ):
            with self.subTest(export=export_name), patch.object(self.workbook, method_name, return_value=sample_bytes):
                exported = load_workbook(BytesIO(getattr(self.workbook, export_name)(departments)), data_only=True)
                self.assertEqual(exported.sheetnames, [
                    "แผนกแรงสูง - รายงาน",
                    "แผนกหม้อแปลง - รายงาน",
                ])


if __name__ == "__main__":
    unittest.main()
